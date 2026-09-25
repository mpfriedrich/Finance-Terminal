"""Excel workbooks from what the app shows — the numbers, not the formatting of a screen.

Every export is raw values (a ratio as 0.153, not "15.3%") with the number format set
in the cell, so Excel can compute with them and a model built on the sheet keeps
working. Each workbook opens with a sheet saying where the numbers came from and as of
when; a spreadsheet that has lost its provenance is how stale figures end up in a model.
"""

from __future__ import annotations

import datetime as dt
import io

import numpy as np
import pandas as pd

# Column-name hints for Excel number formats. First match wins.
_FORMATS = (
    (("margin", "growth", "return_on", "roic", "roe", "roa", "yield", "tax_rate",
      "weight", "vs_", "premium", "_pct", "return"), "0.0%"),
    (("ev_to", "price_to", "_to_ebitda", "coverage", "multiple", "ratio", "beta",
      "z-score", "z_score"), "0.00"),
    (("price", "per_share", "per share"), "#,##0.00"),
    (("value", "cap", "revenue", "income", "ebit", "cash", "debt", "assets",
      "liabilities", "capex", "flow", "book", "amount", "equity"), "#,##0"),
)


def _format_for(column: str) -> str | None:
    name = column.lower()
    for keys, fmt in _FORMATS:
        if any(k in name for k in keys):
            return fmt
    return None


def workbook(sheets: dict[str, pd.DataFrame], title: str, source: str,
             as_of: dt.date | str | None = None) -> bytes:
    """An .xlsx with an 'About' sheet first, then one sheet per frame.

    Sheet names are cut to Excel's 31 characters; empty frames still get their sheet,
    with a line saying so, so a missing table is visible rather than silently absent.
    """
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        about = pd.DataFrame({
            "": ["Title", "Source", "As of", "Exported", "Note"],
            " ": [title, source, str(as_of or "—"),
                  dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
                  "Values are raw numbers with Excel formats; ratios are fractions."],
        })
        about.to_excel(writer, sheet_name="About", index=False)
        for name, frame in sheets.items():
            safe = "".join(c for c in name if c not in "[]:*?/\\")[:31] or "Sheet"
            data = frame if frame is not None and not frame.empty else pd.DataFrame(
                {"note": ["no rows for this table"]})
            data = data.replace([np.inf, -np.inf], np.nan)
            data.to_excel(writer, sheet_name=safe, index=False)
            ws = writer.sheets[safe]
            ws.freeze_panes = "B2"
            header_fill = PatternFill("solid", fgColor="1F1F1F")
            for idx, column in enumerate(data.columns, start=1):
                cell = ws.cell(row=1, column=idx)
                cell.font = Font(bold=True, color="FF9E1B")
                cell.fill = header_fill
                letter = get_column_letter(idx)
                width = max([len(str(column))] + [len(str(v)) for v in data[column].head(200)])
                ws.column_dimensions[letter].width = min(max(width + 2, 8), 48)
                fmt = _format_for(str(column))
                if fmt and pd.api.types.is_numeric_dtype(data[column]):
                    for row in range(2, len(data) + 2):
                        ws.cell(row=row, column=idx).number_format = fmt
        writer.sheets["About"].column_dimensions["A"].width = 12
        writer.sheets["About"].column_dimensions["B"].width = 80
    return buffer.getvalue()


def filename(stem: str, as_of: dt.date | str | None = None) -> str:
    clean = "".join(c if c.isalnum() or c in "-_" else "_" for c in stem).strip("_")
    return f"{clean}_{as_of or dt.date.today()}.xlsx"
