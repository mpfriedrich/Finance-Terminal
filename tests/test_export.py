"""Excel exports: raw values with number formats, provenance first, nothing lost."""
import io

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from portfolio_lab import export


def test_a_workbook_opens_with_its_provenance_and_keeps_raw_values():
    frame = pd.DataFrame({"ticker": ["AAA", "BBB"], "ev_to_ebitda": [6.0, np.inf],
                          "operating_margin": [0.153, 0.2], "market_cap": [1e9, 2e9]})
    data = export.workbook({"Screen": frame}, "Test", "SEC XBRL + prices", "2026-09-24")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["About", "Screen"]
    assert wb["About"]["B2"].value == "Test" and wb["About"]["B4"].value == "2026-09-24"
    ws = wb["Screen"]
    assert ws["C2"].value == 0.153, "a margin stays a fraction, not a string"
    assert ws["C2"].number_format == "0.0%"
    assert ws["B2"].number_format == "0.00"
    assert ws["B3"].value is None, "infinity is written as an empty cell, not an error"
    assert ws["D2"].number_format == "#,##0"


def test_an_empty_table_says_so():
    data = export.workbook({"Comps": pd.DataFrame()}, "T", "S")
    ws = load_workbook(io.BytesIO(data))["Comps"]
    assert ws["A2"].value == "no rows for this table"


def test_sheet_names_are_made_legal():
    data = export.workbook({"A/B: [very long name that Excel refuses]": pd.DataFrame(
        {"x": [1]})}, "T", "S")
    names = load_workbook(io.BytesIO(data)).sheetnames
    assert len(names[1]) <= 31 and "/" not in names[1]


def test_file_names_are_safe():
    assert export.filename("AAPL comps", "2026-09-24") == "AAPL_comps_2026-09-24.xlsx"
